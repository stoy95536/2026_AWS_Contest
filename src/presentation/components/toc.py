"""目錄頁面元件。"""

from pptx.util import Pt, Cm
from pptx.dml.color import RGBColor
from pptx.slide import Slide


def fill_toc(slide: Slide, spec: dict):
    """填充目錄頁。"""
    chapters = [
        "01  市場整體概況",
        "02  同業競爭分析",
        "03  客戶活躍度與獲利能力",
        "04  風險與警訊",
        "05  對台新的策略建議",
    ]
    y = Cm(4.5)
    for ch in chapters:
        txBox = slide.shapes.add_textbox(Cm(3.0), y, Cm(25.0), Cm(1.5))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = f"CHAPTER {ch}"
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
        y += Cm(2.0)
