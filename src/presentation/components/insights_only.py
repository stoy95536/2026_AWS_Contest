"""純洞察文字頁面元件。"""

from pptx.util import Pt, Cm
from pptx.slide import Slide


def fill_insights_only(slide: Slide, spec: dict):
    """只有洞察文字的頁面。"""
    insights = spec.get("insights", [])
    y = Cm(4.5)
    for ins in insights[:6]:
        text = ins.get("text", "") if isinstance(ins, dict) else str(ins)
        txBox = slide.shapes.add_textbox(Cm(1.5), y, Cm(30.0), Cm(1.5))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"• {text}"
        p.font.size = Pt(11)
        y += Cm(1.8)
